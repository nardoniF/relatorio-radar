from __future__ import annotations

import re
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _iter_blocks(text: str, title: str | None = None):
    if title:
        yield True, title.strip()
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = False
        body = line
        if line.startswith("### "):
            body = line[4:]
            heading = True
        elif line.startswith("## "):
            body = line[3:]
            heading = True
        elif line.startswith("# "):
            body = line[2:]
            heading = True
        elif re.match(r"^[A-ZÁÉÍÓÚÃÕÇ0-9][A-ZÁÉÍÓÚÃÕÇ0-9 \-–—\.]{12,}$", line.strip()):
            heading = True
            body = line.strip()
        yield heading, body.strip().strip("*")


def markdown_to_docx(text: str, dest: Path, title: str | None = None) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(8)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)

    for heading, body in _iter_blocks(text):
        if title and heading and body == title.strip():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if heading else WD_ALIGN_PARAGRAPH.JUSTIFY
        if not heading:
            p.paragraph_format.first_line_indent = Cm(1.5)
        run = p.add_run(body)
        run.bold = heading
        run.font.name = "Times New Roman"
        run.font.size = Pt(12 if not heading else 13)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def markdown_to_pdf(text: str, dest: Path, title: str | None = None) -> Path:
    """Gera PDF simples (Times) para protocolar / enviar — sem depender do Word."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4
    margin_l, margin_r, margin_t, margin_b = 70, 50, 60, 60
    width = page.rect.width - margin_l - margin_r
    y = margin_t
    fontsize_body = 11
    fontsize_head = 12
    leading = 1.45

    def new_page():
        nonlocal page, y
        page = doc.new_page(width=595, height=842)
        y = margin_t

    def ensure_space(h: float):
        nonlocal y
        if y + h > page.rect.height - margin_b:
            new_page()

    blocks = list(_iter_blocks(text, title=title))
    for heading, body in blocks:
        size = fontsize_head if heading else fontsize_body
        font = "times-bold" if heading else "times-roman"
        # wrap
        words = body.split()
        lines: list[str] = []
        cur = ""
        for w in words:
            trial = (cur + " " + w).strip()
            tw = fitz.get_text_length(trial, fontname=font, fontsize=size)
            if tw <= width or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        for i, line in enumerate(lines):
            line_h = size * leading
            ensure_space(line_h)
            x = margin_l
            if heading:
                tw = fitz.get_text_length(line, fontname=font, fontsize=size)
                x = margin_l + max(0, (width - tw) / 2)
            page.insert_text((x, y + size), line, fontname=font, fontsize=size, color=(0, 0, 0))
            y += line_h
        y += size * 0.35

    doc.save(dest)
    doc.close()
    return dest


def save_peca(text: str, case_dir: Path, base_name: str, title: str | None = None) -> dict:
    """Grava Word (editar) + PDF (enviar) na pasta do processo."""
    base = base_name
    if base.lower().endswith(".docx"):
        base = base[:-5]
    elif base.lower().endswith(".pdf"):
        base = base[:-4]
    docx_path = case_dir / f"{base}.docx"
    pdf_path = case_dir / f"{base}.pdf"
    markdown_to_docx(text, docx_path, title=title)
    markdown_to_pdf(text, pdf_path, title=title)
    return {"docx": docx_path.name, "pdf": pdf_path.name, "base": base}
